import os
import re
import shutil
import docx

doc_path = r'Nội Bệnh Lý/RHM42 (2018-2019).docx'
backup_path = r'Nội Bệnh Lý/RHM42 (2018-2019)_backup.docx'

if not os.path.exists(backup_path):
    shutil.copyfile(doc_path, backup_path)
    print(f"Backup created at {backup_path}")

answers = {
    1: 'A', 2: 'C', 3: 'C', 4: 'A', 5: 'A',
    6: 'C', 7: 'C', 8: 'C', 9: 'D', 10: 'D',
    11: 'D', 12: 'C', 13: 'C', 14: 'B', 15: 'D',
    16: 'C', 17: 'D', 18: 'D', 19: 'B', 20: 'C',
    21: 'A', 22: 'A', 23: 'A', 24: 'D', 25: 'B',
    26: 'B', 27: 'C', 28: 'C', 29: 'B', 30: 'B',
    31: 'A', 32: 'C', 33: 'B', 34: 'D', 35: 'A',
    36: 'B', 37: 'A', 38: 'C', 39: 'B', 40: 'D',
    41: 'A', 42: 'C', 43: 'B', 44: 'D', 45: 'B',
    46: 'B', 47: 'A', 48: 'A', 49: 'D', 50: 'D',
    51: 'D', 52: 'B', 53: 'A', 54: 'A', 55: 'B',
    56: 'C', 57: 'B', 58: 'A', 59: 'B', 60: 'A'
}

doc = docx.Document(doc_path)

current_q = None
question_count = 0
options_updated = 0
all_questions_found = {}

for p_idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    
    q_match = re.match(r'^Câu\s*(\d+)[:.]', text, re.IGNORECASE)
    if q_match:
        current_q = int(q_match.group(1))
        question_count += 1
        all_questions_found[current_q] = []
        continue
    
    opt_match = re.match(r'^([A-D])\s*[\.:]', text)
    if opt_match and current_q is not None:
        opt_letter = opt_match.group(1).upper()
        all_questions_found[current_q].append((opt_letter, p))
        
        # Remove any existing checkmarks or checkmark characters from runs
        for r in p.runs:
            r.text = r.text.replace('✅', '').replace('✓', '').replace('✔', '')
            
        # If this option is the correct answer according to key, add ' ✅'
        if current_q in answers and opt_letter == answers[current_q]:
            if p.runs and p.runs[-1].text.endswith(' '):
                p.runs[-1].text = p.runs[-1].text.rstrip()
            p.add_run(' ✅')
            options_updated += 1

doc.save(doc_path)
print(f"Processed {len(all_questions_found)} questions.")
print(f"Updated {options_updated} correct answers.")
