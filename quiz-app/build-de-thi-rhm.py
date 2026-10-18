"""Build the RHM past-exam sets (Sâu Răng Học, Phục Hình Tháo Lắp) for the website.

Each exam was transcribed by hand from the scanned PDFs in "Đề Thi/" into
<subject>/nguon/RHM Kxx.txt. This script validates those files and writes, per exam:
  <subject>/RHM Kxx.json   - questions read by build-static.js
  <subject>/RHM Kxx.docx   - the same exam as a Word file
  <subject>/hinh/*.webp    - figures cropped from the PDF (only when the PDF is present)

Source format:
  Câu N. question text
  [hinh: page x0 y0 x1 y1]   optional figure, page of the PDF and its box as page fractions
  *A. option   answer marked in the scanned exam
  @A. option   answer chosen by AI (exam unmarked; taken from another year or reasoned)
  ?A. / !B.    exam marked A but AI picked B (shown with a warning badge)
  > note       explanation (or the AI's reasoning for @ and ?/! answers)
Questions with no marked option have no answer key.
"""

import io
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PDF_DIR = ROOT / "Đề Thi"
SUBJECTS = {
    "Đề thi Sâu Răng Học RHM": {"short": "Sâu răng", "slug": "srh", "pdf": None},
    "Đề thi Phục Hình Tháo Lắp RHM": {
        "short": "Phục hình",
        "slug": "phtl",
        "pdf": "phục hình tháo lắp bán phần nền nhựa ( phục hình I ) - biên soạn Võ Viễn Thông.pdf",
    },
}
QUESTION_RE = re.compile(r"^Câu (\d+)\.\s+(.+)$")
OPTION_RE = re.compile(r"^([*@?!]?)([A-F])\.\s+(.+)$")
IMAGE_RE = re.compile(r"^\[hinh: (\d+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)\]$")


def fail(path, message):
    raise SystemExit(f"{path.name}: {message}")


def parse(path):
    header, items = [], []
    current = None
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("#") and not items:
            header.append(line.lstrip("# ").strip())
            continue
        match = QUESTION_RE.match(line)
        if match:
            current = {"number": int(match.group(1)), "question": match.group(2).strip(),
                       "options": [], "marks": {}, "notes": [], "image": None}
            items.append(current)
            continue
        if current is None:
            fail(path, f"text before the first question: {line}")
        if line.startswith("> "):
            current["notes"].append(line[2:].strip())
            continue
        match = IMAGE_RE.match(line)
        if match:
            if current["options"] or current["image"]:
                fail(path, f"misplaced figure in question {current['number']}")
            page, *box = match.groups()
            current["image"] = {"page": int(page), "box": [float(v) for v in box]}
            continue
        match = OPTION_RE.match(line)
        if not match or current["notes"]:
            fail(path, f"unexpected line in question {current['number']}: {line}")
        mark, label, text = match.groups()
        if label != "ABCDEF"[len(current["options"])]:
            fail(path, f"option order differs in question {current['number']}")
        current["options"].append({"label": label, "text": text.strip()})
        if mark:
            current["marks"].setdefault(mark, []).append(len(current["options"]) - 1)
    return header, items


def resolve(path, item):
    """Turn option markers into the website's answer fields."""
    marks, note = item["marks"], " ".join(item["notes"]) or None
    if len(item["options"]) < 2:
        fail(path, f"question {item['number']} has fewer than two options")
    if any(len(v) > 1 for v in marks.values()) or set(marks) not in (set(), {"*"}, {"@"}, {"?", "!"}):
        fail(path, f"question {item['number']} has inconsistent answer marks {marks}")
    out = {"correctIndex": None, "explanation": None, "ai": False, "aiCheck": False, "aiNote": None}
    if "*" in marks:
        out.update(correctIndex=marks["*"][0], explanation=note)
    elif "@" in marks:
        if not note:
            fail(path, f"question {item['number']} has an AI answer without a reason")
        out.update(correctIndex=marks["@"][0], ai=True, aiNote=note)
    elif "!" in marks:
        if not note:
            fail(path, f"question {item['number']} overrides the exam without a reason")
        out.update(correctIndex=marks["!"][0], ai=True, aiCheck=True, aiNote=note)
    else:
        if not note:
            fail(path, f"question {item['number']} has no answer and no note saying why")
        out["explanation"] = note
    return out


def crop_figure(pdf, figure, target):
    import pymupdf
    from PIL import Image

    page = pdf[figure["page"] - 1]
    x0, y0, x1, y1 = figure["box"]
    rect = page.rect
    clip = pymupdf.Rect(rect.width * x0, rect.height * y0, rect.width * x1, rect.height * y1)
    zoom = 520 / clip.width
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=clip)
    image = Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("L")
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target, "WEBP", quality=82)


def create_docx(title, header, questions, folder, target):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(1.5)

    heading = doc.add_paragraph()
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    for line in header[1:]:
        doc.add_paragraph(line).runs[0].italic = True

    for q in questions:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(7)
        para.paragraph_format.keep_with_next = True
        para.add_run(f"Câu {q['number']}. ").bold = True
        para.add_run(q["question"])
        if q["image"]:
            from PIL import Image

            buffer = io.BytesIO()
            Image.open(folder / q["image"]["file"]).save(buffer, "PNG")
            buffer.seek(0)
            doc.add_picture(buffer, width=Inches(2.2))
        for option in q["options"]:
            line = doc.add_paragraph()
            line.paragraph_format.left_indent = Inches(0.18)
            line.add_run(f"{option['label']}. ").bold = True
            line.add_run(option["text"])
        index = q["correctIndex"]
        answer = "chưa có" if index is None else q["options"][index]["label"]
        if q["aiCheck"]:
            answer += " (AI làm đáp án, khác phương án khoanh trong đề — nên kiểm tra)"
        elif q["ai"]:
            answer += " (AI làm đáp án)"
        doc.add_paragraph().add_run(f"Đáp án: {answer}").bold = True
        note = q["aiNote"] or q["explanation"]
        if note:
            para = doc.add_paragraph()
            para.add_run("Ghi chú: ").bold = True
            para.add_run(note)
    doc.save(target)


def build_subject(subject, info):
    folder = ROOT / subject
    pdf = None
    if info["pdf"] and (PDF_DIR / info["pdf"]).exists():
        import pymupdf

        pdf = pymupdf.open(PDF_DIR / info["pdf"])
    report = []
    for source in sorted((folder / "nguon").glob("RHM K*.txt")):
        title = source.stem
        header, items = parse(source)
        if not header or header[0] != title:
            fail(source, "first header line must be the exam title")
        numbers = [q["number"] for q in items]
        if numbers != sorted(set(numbers)):
            fail(source, "question numbers are repeated or out of order")
        seen = {}
        questions = []
        for item in items:
            key = (item["question"].casefold(), tuple(o["text"].casefold() for o in item["options"]))
            if key in seen and not item["image"]:
                fail(source, f"question {item['number']} duplicates question {seen[key]}")
            seen[key] = item["number"]
            text = item["question"] + "".join(o["text"] for o in item["options"]) + "".join(item["notes"])
            if "�" in text:
                fail(source, f"broken character in question {item['number']}")
            image = None
            if item["image"]:
                name = f"{info['slug']}-{title.split()[-1].lower()}-c{item['number']:02d}.webp"
                image = {"file": f"hinh/{name}", **item["image"]}
                if pdf is not None:
                    crop_figure(pdf, item["image"], folder / image["file"])
                elif not (folder / image["file"]).exists():
                    fail(source, f"figure for question {item['number']} is missing and the PDF is not here")
            questions.append({
                "number": item["number"],
                "question": item["question"],
                "options": item["options"],
                "image": image,
                **resolve(source, item),
            })

        (folder / f"{title}.json").write_text(json.dumps({
            "title": title,
            "subject": subject,
            "source": f"{info['short']} {title}",
            "header": header[1:],
            "questions": questions,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        create_docx(title, header, questions, folder, folder / f"{title}.docx")
        report.append(
            f"  {subject} / {title}: {len(questions)} câu, "
            f"{sum(q['correctIndex'] is not None and not q['ai'] for q in questions)} theo đề, "
            f"{sum(q['ai'] and not q['aiCheck'] for q in questions)} AI, "
            f"{sum(q['aiCheck'] for q in questions)} AI sửa, "
            f"{sum(q['correctIndex'] is None for q in questions)} chưa có đáp án, "
            f"{sum(q['image'] is not None for q in questions)} hình"
        )
    return report


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    lines = []
    for subject, info in SUBJECTS.items():
        lines += build_subject(subject, info)
    print("\n".join(lines))


if __name__ == "__main__":
    sys.exit(main())
