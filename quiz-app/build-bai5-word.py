"""Validate the supplied Bài 5 Word source and build the website data and quiz DOCX."""

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
FOLDER = ROOT / "Sâu Răng Học Azota Thông Võ"
SOURCE = FOLDER / "bai5-word-source.json"
VALIDATED = FOLDER / "bai5-validated.json"
QA = FOLDER / "bai5-qa.json"
OUTPUT = FOLDER / "Bài 5.docx"
QUESTION_RE = re.compile(r"^Câu\s+(\d+)\.\s*(.+)$", re.I)
OPTION_RE = re.compile(r"^(\*?)([A-D])\.\s*(.+)$")
ANSWER_RE = re.compile(r"^Đáp án:\s*([A-D])$")


def plain(value):
    return re.sub(r"\s+", " ", value).strip()


def parse_source(source):
    if source["mediaCount"] or source["tableCount"] or source["inlineShapeCount"]:
        raise ValueError("Source has media or tables requiring manual placement")

    paragraphs = source["paragraphs"]
    items = []
    seen = {}
    index = 0
    while index < len(paragraphs):
        match = QUESTION_RE.fullmatch(paragraphs[index])
        if not match:
            raise ValueError(f"Unexpected source paragraph {index + 1}: {paragraphs[index]}")
        number = int(match.group(1))
        if number != len(items) + 1:
            raise ValueError(f"Missing, repeated, or reordered question {number}")
        question = plain(match.group(2))
        index += 1

        options = []
        marked = []
        for expected in "ABCD":
            if index >= len(paragraphs):
                raise ValueError(f"Question {number} is missing option {expected}")
            option = OPTION_RE.fullmatch(paragraphs[index])
            if not option or option.group(2) != expected:
                raise ValueError(f"Question {number} option order differs at {expected}")
            options.append({"label": expected, "text": plain(option.group(3))})
            if option.group(1):
                marked.append(expected)
            index += 1

        if index >= len(paragraphs):
            raise ValueError(f"Question {number} has no answer line")
        answer_match = ANSWER_RE.fullmatch(paragraphs[index])
        if not answer_match:
            raise ValueError(f"Question {number} has an invalid answer line")
        answer = answer_match.group(1)
        if marked != [answer]:
            raise ValueError(f"Question {number} answer and marked option disagree")
        index += 1

        explanation = None
        if index < len(paragraphs) and not QUESTION_RE.fullmatch(paragraphs[index]):
            if number != 117 or not paragraphs[index].startswith("(Chú ý:"):
                raise ValueError(f"Unclassified paragraph after question {number}")
            explanation = plain(paragraphs[index])
            index += 1

        text = question + " ".join(o["text"] for o in options) + (explanation or "")
        if "\ufffd" in text or any(0x80 <= ord(c) <= 0x9F for c in text):
            raise ValueError(f"Broken character at question {number}")
        key = (question.casefold(), tuple(o["text"].casefold() for o in options))
        if key in seen:
            raise ValueError(f"Question {number} duplicates question {seen[key]}")
        seen[key] = number

        items.append({
            "number": number,
            "question": question,
            "options": options,
            "answer": answer,
            "explanation": explanation,
            "images": [],
            "tables": [],
        })

    if len(items) != 500:
        raise ValueError(f"Expected 500 source questions, got {len(items)}")
    if [q["number"] for q in items] != list(range(1, 501)):
        raise ValueError("Question sequence is incomplete")
    if [q["number"] for q in items if q["explanation"]] != [117]:
        raise ValueError("The note after question 117 was not preserved correctly")
    return items


def create_docx(items):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(1.5)

    title = doc.add_paragraph(style="Title")
    title.add_run("Bài 5")
    title.style.font.name = "Arial"
    title.style.font.size = Pt(18)
    title.style.font.bold = True
    title.style.font.color.rgb = None
    title_properties = title.style._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    title.paragraph_format.space_after = Pt(12)

    for item in items:
        question = doc.add_paragraph()
        question.paragraph_format.space_before = Pt(7)
        question.paragraph_format.space_after = Pt(3)
        question.paragraph_format.keep_with_next = True
        question.add_run(f"Câu {item['number']}. ").bold = True
        question.add_run(item["question"])
        for option in item["options"]:
            line = doc.add_paragraph()
            line.paragraph_format.left_indent = Inches(0.18)
            line.add_run(f"{option['label']}. ").bold = True
            line.add_run(option["text"])
        answer = doc.add_paragraph()
        answer.add_run(f"Đáp án: {item['answer']}").bold = True
        answer.paragraph_format.space_after = Pt(7)
        if item["explanation"]:
            explanation = doc.add_paragraph()
            explanation.add_run("Giải thích: ").bold = True
            explanation.add_run(item["explanation"])
    doc.save(OUTPUT)


def main():
    source = json.loads(SOURCE.read_text(encoding="utf-8"))
    items = parse_source(source)
    report = {
        "sourceFile": source["sourceFile"],
        "sourceSha256": source["sourceSha256"],
        "answerSource": "Answer lines and matching starred options in supplied Word file",
        "sourceQuestionCount": len(items),
        "extractedCount": len(items),
        "uniqueCount": len(items),
        "wordCount": len(items),
        "websiteCount": len(items),
        "missingNumbers": [],
        "duplicateCount": 0,
        "withoutAnswer": [],
        "withoutExplanation": sum(q["explanation"] is None for q in items),
        "explanationQuestions": [q["number"] for q in items if q["explanation"]],
        "imageQuestions": 0,
        "tableQuestions": 0,
    }
    VALIDATED.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    QA.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    create_docx(items)
    print(f"Created Bài 5.docx with {len(items)} questions")


if __name__ == "__main__":
    main()
