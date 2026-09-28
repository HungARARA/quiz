"""Create the reviewable Word source for the Azota Bài 1 quiz."""

import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent.parent
FOLDER = ROOT / "Sâu Răng Học Azota Thông Võ"
SOURCE = FOLDER / "azota-source.json"
VALIDATED = FOLDER / "azota-validated.json"
OUTPUT = FOLDER / "Bài 1.docx"


def validate(items):
    if len(items) != 500:
        raise ValueError(f"Azota source contains {len(items)} questions, expected 500")
    clean = []
    seen = set()
    for expected, raw in enumerate(items, 1):
        q = dict(raw)
        if q["number"] != expected:
            raise ValueError(f"Question sequence differs at {expected}: {q['number']}")
        if not q["question"].strip():
            raise ValueError(f"Question {expected} has no text")
        options = [dict(o) for o in q["options"]]
        if [o["label"] for o in options] != ["A", "B", "C", "D"]:
            raise ValueError(f"Question {expected} option order differs from Azota")
        if any(not o["text"].strip() for o in options):
            raise ValueError(f"Question {expected} contains an empty option")
        if q["answer"] is None:
            # Azota's last option for question 200 ends in "Đáp án: C\nCâu".
            # Recover only the letter explicitly present in its preserved HTML.
            match = re.search(r"Đáp án\s*:\s*([A-D])(?:<br>|<)", options[-1]["html"])
            if not match:
                raise ValueError(f"Question {expected} has no explicit Azota answer")
            q["answer"] = match.group(1)
            options[-1]["text"] = re.sub(
                r"\s*Đáp án\s*:\s*[A-D]\s*(?:Câu)?\s*$", "", options[-1]["text"]
            ).strip()
        if q["answer"] not in [o["label"] for o in options]:
            raise ValueError(f"Question {expected} has an invalid answer")
        if q["images"] or q["tables"]:
            raise ValueError(f"Question {expected} has media requiring placement")
        key = (q["question"].strip(), tuple(o["text"].strip() for o in options))
        if key in seen:
            raise ValueError(f"Question {expected} duplicates another question")
        seen.add(key)
        q["options"] = options
        clean.append(q)
    return clean


def create_docx(items):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(1.5)

    title = doc.add_paragraph(style="Title")
    title.add_run("Bài 1")
    title.style.font.name = "Arial"
    title.style.font.size = Pt(18)
    title.style.font.bold = True
    title.style.font.color.rgb = None
    title_properties = title.style._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    title.paragraph_format.space_after = Pt(12)

    for q in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"Câu {q['number']}. ").bold = True
        p.add_run(q["question"])
        for option in q["options"]:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Inches(0.18)
            op.add_run(f"{option['label']}. ").bold = True
            op.add_run(option["text"])
        answer = doc.add_paragraph()
        answer.add_run(f"Đáp án: {q['answer']}").bold = True
        answer.paragraph_format.space_after = Pt(7)
        if q["explanation"]:
            exp = doc.add_paragraph()
            exp.add_run("Giải thích: ").bold = True
            exp.add_run(q["explanation"])
            exp.paragraph_format.space_after = Pt(7)

    doc.save(OUTPUT)


def main():
    items = json.loads(SOURCE.read_text(encoding="utf-8"))
    items = validate(items)
    VALIDATED.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    create_docx(items)
    print(f"Created {OUTPUT.name} with {len(items)} questions")


if __name__ == "__main__":
    main()
