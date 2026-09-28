"""Validate the Azota Bài 2 extraction and create its Word study copy."""

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
FOLDER = ROOT / "Sâu Răng Học Azota Thông Võ"
SOURCE = FOLDER / "azota-bai2-source.json"
ANSWER_KEY = FOLDER / "azota-bai2-answer-key.json"
VALIDATED = FOLDER / "azota-bai2-validated.json"
QA = FOLDER / "azota-bai2-qa.json"
OUTPUT = FOLDER / "Bài 2.docx"
FORMULA = "$Ca_{10}(PO_{4})_{6}(OH)_{2}$"
FORMULA_UNICODE = "Ca₁₀(PO₄)₆(OH)₂"


def validate(raw_items, answer_data):
    if len(raw_items) != 500:
        raise ValueError(f"Azota lists 500 questions; extracted {len(raw_items)}")
    if answer_data["sourceUrl"] != "https://azota.vn/vi/de-thi/xol3ym":
        raise ValueError("Answer key belongs to another Azota exam")
    answer_rows = answer_data["answers"]
    if len(answer_rows) != 500:
        raise ValueError(f"Expected 500 Azota answer rows, found {len(answer_rows)}")

    items = []
    seen = {}
    duplicates = []
    removed_pdf = []
    for expected, raw in enumerate(raw_items, 1):
        q = dict(raw)
        if q["number"] != expected:
            raise ValueError(f"Missing or out-of-order question at {expected}")
        if not q["question"].strip():
            raise ValueError(f"Question {expected} is empty")
        if [o["label"] for o in q["options"]] != ["A", "B", "C", "D"]:
            raise ValueError(f"Option order differs at question {expected}")
        if q["answer"] is not None:
            raise ValueError(f"Unexpected answer in original question page at {expected}")
        key_row = answer_rows[expected - 1]
        if (
            key_row["number"] != expected
            or key_row["id"] != str(q["azotaId"])
            or key_row["answer"] not in [o["label"] for o in q["options"]]
        ):
            raise ValueError(f"Answer key does not match Azota question {expected}")
        q["answer"] = key_row["answer"]
        if q["explanation"] is not None:
            raise ValueError(f"Unexpected explanation supplied for question {expected}")
        if q["images"] or q["tables"]:
            raise ValueError(f"Question {expected} has media requiring placement")

        q["question"] = q["question"].replace(FORMULA, FORMULA_UNICODE)
        options = [dict(o) for o in q["options"]]
        for o in options:
            if o["text"].endswith("\nPDF"):
                o["text"] = o["text"][:-4].strip()
                removed_pdf.append(expected)
            if not o["text"].strip():
                raise ValueError(f"Empty option at question {expected}")
        q["options"] = options

        text = q["question"] + " ".join(o["text"] for o in options)
        if "\ufffd" in text:
            raise ValueError(f"Replacement character at question {expected}")
        key = (q["question"], tuple(o["text"] for o in options))
        if key in seen:
            if q["answer"] != answer_rows[seen[key] - 1]["answer"]:
                raise ValueError(f"Repeated question {expected} has a different answer")
            duplicates.append({"number": expected, "duplicates": seen[key]})
            continue
        seen[key] = expected
        items.append(q)

    if len(duplicates) != 50 or [d["number"] for d in duplicates] != list(range(51, 101)):
        raise ValueError("Duplicate pattern differs from Azota's questions 51–100")
    if removed_pdf != list(range(51, 100)):
        raise ValueError("Unexpected placement of Azota's trailing PDF marker")
    if len(items) != 450:
        raise ValueError(f"Expected 450 distinct Azota questions, found {len(items)}")
    report = {
        "sourceUrl": "https://azota.vn/vi/de-thi/xol3ym",
        "azotaCount": 500,
        "extractedCount": len(raw_items),
        "azotaAnswerCount": len(answer_rows),
        "answerSource": answer_data["answerSource"],
        "uniqueCount": len(items),
        "wordCount": len(items),
        "websiteCount": len(items),
        "missingNumbers": [],
        "duplicateCount": len(duplicates),
        "duplicatePairs": duplicates,
        "removedDuplicateNumbers": [d["number"] for d in duplicates],
        "withoutAnswer": [],
        "withoutExplanation": len(items),
        "imageQuestions": 0,
        "tableQuestions": 0,
        "trailingPdfMarkerRemovedFromOptions": removed_pdf,
        "formulaRenderedAsUnicodeAt": [102],
    }
    return items, report


def create_docx(items):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(1.5)

    title = doc.add_paragraph(style="Title")
    title.add_run("Bài 2")
    title.style.font.name = "Arial"
    title.style.font.size = Pt(18)
    title.style.font.bold = True
    title.style.font.color.rgb = None
    title_properties = title.style._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    title.paragraph_format.space_after = Pt(12)

    for q in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"Câu {q['number']}. ").bold = True
        p.add_run(q["question"])
        for option in q["options"]:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Inches(0.18)
            op.add_run(f"{option['label']}. ").bold = True
            op.add_run(option["text"])
        answer = doc.add_paragraph()
        answer.add_run(f"Đáp án: {q['answer']}").bold = True
        answer.paragraph_format.space_after = Pt(7)

    doc.save(OUTPUT)


def main():
    raw_items = json.loads(SOURCE.read_text(encoding="utf-8"))
    answer_data = json.loads(ANSWER_KEY.read_text(encoding="utf-8"))
    items, report = validate(raw_items, answer_data)
    VALIDATED.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    QA.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    create_docx(items)
    print(f"Created Bài 2.docx with {len(items)} questions")


if __name__ == "__main__":
    main()
