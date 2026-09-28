"""Validate Azota Bài 3 and create the matching Word quiz."""

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
FOLDER = ROOT / "Sâu Răng Học Azota Thông Võ"
SOURCE = FOLDER / "azota-bai3-source.json"
VALIDATED = FOLDER / "azota-bai3-validated.json"
QA = FOLDER / "azota-bai3-qa.json"
OUTPUT = FOLDER / "Bài 3.docx"
SOURCE_URL = "https://azota.vn/vi/de-thi/sqdhqn"


def plain(text):
    return re.sub(r"\s+", " ", text).strip()


def validate(rows):
    if len(rows) != 500:
        raise ValueError(f"Azota lists 500 questions; extracted {len(rows)}")

    items = []
    seen = {}
    for expected, raw in enumerate(rows, 1):
        if raw["number"] != expected:
            raise ValueError(f"Missing or out-of-order question at {expected}")
        if not raw["azotaId"] or not raw["question"].strip():
            raise ValueError(f"Missing ID or question text at {expected}")
        if [o["label"] for o in raw["options"]] != ["A", "B", "C", "D"]:
            raise ValueError(f"Option order differs at question {expected}")
        match = re.fullmatch(r"Đáp án đúng:\s*([A-D])", raw["answerRaw"].strip())
        if not match:
            raise ValueError(f"Missing or invalid Azota answer at question {expected}")
        if raw["images"] or raw["tables"] or "<img" in raw["questionHtml"].lower():
            raise ValueError(f"Question {expected} has media requiring placement")
        if any("<img" in option["html"].lower() for option in raw["options"]):
            raise ValueError(f"Option image at question {expected} requires placement")

        question = plain(raw["question"])
        options = [{"label": o["label"], "text": plain(o["text"])} for o in raw["options"]]
        if any(not option["text"] for option in options):
            raise ValueError(f"Empty option at question {expected}")
        if expected == 300:
            # Azota renders an over-arrow annotation as scattered text on extraction.
            # Apply only the user-approved faithful transcription here.
            original = options[2]["text"]
            expected_fragment = "Chụp X → 𝑞 𝑢 𝑎 𝑛 𝑔 𝑝 𝑎 𝑛 𝑜 𝑟 𝑎 𝑚 𝑎 Đếm số răng còn lại → Làm răng giả tháo lắp."
            if original != expected_fragment:
                raise ValueError("Azota question 300 math source changed; review before publishing")
            options[2]["text"] = "Chụp X —(quang panorama)→ Đếm số răng còn lại → Làm răng giả tháo lắp."
        explanation = plain(raw["explanation"]) or None
        text = question + " ".join(o["text"] for o in options)
        if "\ufffd" in text or any(ord(char) in range(0x80, 0xA0) for char in text):
            raise ValueError(f"Broken character at question {expected}")

        duplicate_key = (question.casefold(), tuple(o["text"].casefold() for o in options))
        if duplicate_key in seen:
            raise ValueError(f"Duplicate question {expected} matches {seen[duplicate_key]}")
        seen[duplicate_key] = expected

        items.append({
            "number": expected,
            "azotaId": raw["azotaId"],
            "question": question,
            "options": options,
            "answer": match.group(1),
            "explanation": explanation,
            "images": raw["images"],
            "tables": raw["tables"],
        })

    if len({q["azotaId"] for q in items}) != 500:
        raise ValueError("Azota question IDs are repeated")
    report = {
        "sourceUrl": SOURCE_URL,
        "answerSource": "Azota submitted-attempt result page, Đáp án đúng field",
        "azotaCount": 500,
        "extractedCount": len(rows),
        "uniqueCount": len(items),
        "wordCount": len(items),
        "websiteCount": len(items),
        "missingNumbers": [],
        "duplicateCount": 0,
        "withoutAnswer": [],
        "withoutExplanation": sum(q["explanation"] is None for q in items),
        "imageQuestions": sum(bool(q["images"]) for q in items),
        "tableQuestions": sum(bool(q["tables"]) for q in items),
        "mathTranscribedAt": [300],
    }
    return items, report


def create_docx(items):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(1.5)

    title = doc.add_paragraph(style="Title")
    title.add_run("Bài 3")
    title.style.font.name = "Arial"
    title.style.font.size = Pt(18)
    title.style.font.bold = True
    title.style.font.color.rgb = None
    title_properties = title.style._element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    title.paragraph_format.space_after = Pt(12)

    for q in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"Câu {q['number']}. ").bold = True
        p.add_run(q["question"])
        for option in q["options"]:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Inches(0.18)
            op.add_run(f"{option['label']}. ").bold = True
            op.add_run(option["text"])
        answer = doc.add_paragraph()
        answer.add_run(f"Đáp án: {q['answer']}").bold = True
        if q["explanation"]:
            explanation = doc.add_paragraph()
            explanation.add_run("Giải thích: ").bold = True
            explanation.add_run(q["explanation"])
        answer.paragraph_format.space_after = Pt(7)
    doc.save(OUTPUT)


def main():
    rows = json.loads(SOURCE.read_text(encoding="utf-8"))
    items, report = validate(rows)
    VALIDATED.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    QA.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    create_docx(items)
    print(f"Created Bài 3.docx with {len(items)} questions")


if __name__ == "__main__":
    main()
